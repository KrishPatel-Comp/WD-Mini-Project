import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Set page margins
sections = doc.sections
for s in sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.85)
    s.right_margin = Inches(0.85)

# Color constants
PRIMARY = RGBColor(39, 98, 235)     # #2762eb
NAVY = RGBColor(24, 34, 60)         # #18223c
MUTED = RGBColor(90, 105, 130)      # #5a6982
DARK = RGBColor(25, 35, 60)         # #19233c
TEAL = RGBColor(0, 168, 132)        # #00a884

def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_banner(title, subtitle):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    set_cell_background(c, "18223C")
    set_cell_margins(c, top=260, bottom=260, left=260, right=260)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r_tag = p.add_run("UNIBRIDGE · WEB DEVELOPMENT PROJECT REPORT\n")
    r_tag.font.size = Pt(10)
    r_tag.font.bold = True
    r_tag.font.color.rgb = RGBColor(94, 161, 255)
    
    r_title = p.add_run(f"{title}\n")
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(255, 255, 255)
    
    r_sub = p.add_run(subtitle)
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(200, 215, 240)
    doc.add_paragraph()

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = NAVY

def add_heading_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = PRIMARY

def add_p(text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.18
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.bold = True
        r_b.font.color.rgb = DARK
        r_b.font.size = Pt(10)
    r = p.add_run(text)
    r.font.color.rgb = DARK
    r.font.size = Pt(10)
    return p

def add_bullet(bold_title, desc):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    rb = p.add_run(bold_title + ": ")
    rb.font.bold = True
    rb.font.color.rgb = DARK
    rb.font.size = Pt(10)
    rt = p.add_run(desc)
    rt.font.color.rgb = DARK
    rt.font.size = Pt(10)

def add_image_with_caption(img_path, caption):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(f"Figure: {caption}")
        r_cap.font.size = Pt(9)
        r_cap.font.italic = True
        r_cap.font.color.rgb = MUTED

# 1. Header Banner
add_header_banner(
    "UniBridge Student Portal",
    "Modular University Admissions, Application Lifecycle Tracking & Simulated Fee Payments"
)

# Meta Summary Table
t_meta = doc.add_table(rows=4, cols=2)
t_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ("Project Name", "UniBridge Student Onboarding & Admissions Portal"),
    ("GitHub Repository", "https://github.com/KrishPatel-Comp/WD-Mini-Project.git"),
    ("Commit ID", "41c6d53 (Branch: main) — Full UI Overhaul & Module Integration"),
    ("Tech Stack", "PHP 8.2 (Apache), MariaDB 11.4 / MySQL, Vanilla CSS 3 (EduAdmin System), JS ES Modules, Docker")
]
for idx, (k, v) in enumerate(meta_data):
    c0 = t_meta.rows[idx].cells[0]
    c1 = t_meta.rows[idx].cells[1]
    set_cell_background(c0, "F0F4FA")
    set_cell_margins(c0, 60, 60, 100, 100)
    set_cell_margins(c1, 60, 60, 100, 100)
    p0 = c0.paragraphs[0]
    p0.add_run(k).font.bold = True
    p0.runs[0].font.size = Pt(9.5)
    p0.runs[0].font.color.rgb = NAVY
    p1 = c1.paragraphs[0]
    r1 = p1.add_run(v)
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = DARK
    if "http" in v:
        r1.font.color.rgb = PRIMARY
doc.add_paragraph()

# ── 1. Executive Summary ──────────────────────────────────────────
add_heading_1("1. Executive Summary")
add_p("UniBridge is a centralized, deployment-ready higher education student onboarding web application. Built for colleges and universities, the portal bridges the gap between student applicants, academic faculty, and bursar offices. During this engineering iteration, the entire platform underwent a comprehensive architectural upgrade: fixing UI overlapping issues across mobile and desktop breakpoints, implementing the modern EduAdmin theme, introducing end-to-end Application Lifecycle Tracking (Draft to Decision Timeline), and deploying a sandboxed multi-method Fee Payment & Receipt System.")

# ── 2. UI/UX Design System & Anti-Overlap Engineering ─────────────
add_heading_1("2. UI/UX Design System & Anti-Overlap Engineering")
add_p("Prior iterations exhibited component crowding and text truncation issues on mid-sized laptop and mobile viewports. The interface was re-architected with a curated Vanilla CSS design system that ensures crisp aesthetics, dynamic micro-animations, and fluid responsiveness without external framework bloat:")

add_bullet("EduAdmin Navy Visual Language", "Crafted with dark navy sidebars (#18223c), deep sapphire interactive states (#2762eb), and clean light canvas cards (#f0f4fa) with subtle drop shadows.")
add_bullet("Responsive Breakpoint Grid", "Custom CSS media queries at 1280px, 1024px, 992px, 768px, and 480px dynamically collapse 4-column course grids and 3-column widget rows to prevent layout breakage.")
add_bullet("Mobile Drawer Navigation", "Integrated a responsive hamburger navigation toggle with backdrop blur overlay, ensuring intuitive navigation on handheld touchscreens.")
add_bullet("Typography & Aesthetics", "Imported Google Fonts Inter and Manrope for clean hierarchy, high legibility, and modern academic SaaS polish.")

# ── 3. Module Overview & Application Walkthrough ───────────────────
add_heading_1("3. Key System Modules & Live Application Walkthrough")

# Module A: Auth
add_heading_2("3.1 Student Authentication & Account Provisioning")
add_p("Secure authentication with bcrypt password hashing (PASSWORD_DEFAULT), CSRF protection tokens, session regeneration, and automated student identification generation (e.g. UB2026-2136).")
add_image_with_caption("c:/WD project/docs/screenshots/01_login_page.png", "EduAdmin Branded Student Login Interface")
add_image_with_caption("c:/WD project/docs/screenshots/02_signup_page.png", "Student Registration & Automatic Admissions ID Provisioning")

# Module B: Dashboard
add_heading_2("3.2 Student Overview Dashboard")
add_p("The dashboard serves as the applicant mission control, calculating onboarding readiness scores, displaying active enrolled courses, SVG activity breakdown charts, and immediate action items.")
add_image_with_caption("c:/WD project/docs/screenshots/03_dashboard_page.png", "Student Overview Dashboard with Readiness Meter and Course Cards")

# Module C: Tracking
add_heading_2("3.3 Feature 5: Application Lifecycle & Status Tracking")
add_p("The application lifecycle tracker provides transparency across the entire review process:")
add_bullet("8 Lifecycle States", "Draft, Submitted, Under Review, Documents Pending, Correction Required, Waitlisted, Approved, and Rejected.")
add_bullet("Milestone Status Timeline", "Chronological timeline with completed step checkmarks, active pulsating stage badges, and official committee remarks.")
add_bullet("Status Stats Cards", "Four high-level KPI cards displaying Application ID, Lifecycle State, Submission Date, and Document Compliance.")
add_bullet("Admissions Review Simulator", "An embedded interactive testing tool enabling faculty or evaluators to simulate admissions transitions and trigger real-time notification alerts.")
add_bullet("Real-Time Notifications Feed", "Alert feed reflecting all status transitions with unread counters and severity indicators (Success, Warning, Alert, Info).")
add_image_with_caption("c:/WD project/docs/screenshots/04_application_tracking.png", "Application Milestones Timeline, Document Compliance & Review Simulator")

# Module D: Payments
add_heading_2("3.4 Feature 6: Fee Payment, StroWallet Checkout & Receipts")
add_p("A robust simulated financial module built for student projects and production deployments:")
add_bullet("Two-Stage Fee Architecture", "Stage 1 Application Processing Fee (NGN 300.00) and Stage 2 Enrolment & Seat Confirmation Fee (NGN 700.00).")
add_bullet("Simulated Payment Gateway Modal", "StroWallet-powered sandboxed modal checkout supporting Card, Balance, and Direct Bank Transfer payments.")
add_bullet("Transaction History Table", "Searchable real-time ledger tracking references, paid amounts, timestamps, and payment statuses (Paid/Pending).")
add_bullet("Official Electronic Receipt", "Printable and downloadable PDF-ready receipt modal featuring verified student credentials, course metadata, and authenticated registrar seals.")
add_image_with_caption("c:/WD project/docs/screenshots/05_fee_payments.png", "Fee Payment Centre with Two-Stage Fee Cards & StroWallet Transaction History")
add_image_with_caption("c:/WD project/docs/screenshots/06_electronic_receipt.png", "Authenticated Electronic Fee Receipt with window.print() Support")

# ── 4. Technical Architecture & Database Design ───────────────────
add_heading_1("4. Technical Architecture & Database Schema")
add_p("UniBridge utilizes a decoupled model with PHP PDO backends serving normalized JSON REST APIs to vanilla ES module clients. Automatic table bootstrapping is performed on startup:")

t_db = doc.add_table(rows=6, cols=3)
t_db.alignment = WD_TABLE_ALIGNMENT.CENTER
db_headers = ["Table Name", "Primary Purpose", "Key Columns / Enums"]
for col_idx, h in enumerate(db_headers):
    c = t_db.rows[0].cells[col_idx]
    set_cell_background(c, "18223C")
    set_cell_margins(c, 80, 80, 100, 100)
    p = c.paragraphs[0]
    p.add_run(h).font.bold = True
    p.runs[0].font.size = Pt(9.5)
    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

schema_rows = [
    ("student_login", "Authentication credentials & audit timestamps", "student_id (PK), admission_no, email, password_hash, last_login"),
    ("student_registration", "Student profile & program selection", "student_id (FK), full_name, course, phone, admission_year, address"),
    ("application_track", "Lifecycle milestone status & committee notes", "application_id (PK), student_id, status (ENUM), status_note, submitted_at"),
    ("fee_payments", "Financial transaction records & verification", "payment_id (PK), student_id, reference, fee_type, amount, status, paid_at"),
    ("student_notifications", "Push notifications triggered on status events", "notification_id (PK), student_id, title, message, type (ENUM), is_read")
]

for row_idx, data in enumerate(schema_rows, start=1):
    for col_idx, text in enumerate(data):
        c = t_db.rows[row_idx].cells[col_idx]
        set_cell_background(c, "FAFCFF" if row_idx % 2 == 1 else "FFFFFF")
        set_cell_margins(c, 60, 60, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(text)
        r.font.size = Pt(9)
        r.font.color.rgb = DARK
        if col_idx == 0:
            r.font.bold = True
doc.add_paragraph()

# ── 5. Dockerization & Deployment ─────────────────────────────────
add_heading_1("5. Containerization & Deployment Pipeline")
add_p("The application is engineered for zero-dependency containerized deployments:")
add_bullet("Dockerfile", "Based on official php:8.2-apache, configuring pdo_mysql, mod_rewrite, custom php.ini configurations (10MB upload limits, disabled display_errors, enabled error_logging), and Apache healthchecks.")
add_bullet("docker-compose.yml", "Orchestrates three services: the PHP web app (port 8080), MariaDB 11.4 database (port 3306 with healthcheck dependencies), and phpMyAdmin (port 8081).")
add_bullet("Environment Configuration", "Parameterized database connection strings via DB_HOST, DB_NAME, DB_USER, and DB_PASS for seamless CI/CD and production environments.")

# ── 6. Conclusion ──────────────────────────────────────────────────
add_heading_1("6. Conclusion")
add_p("The UniBridge portal upgrade fulfills all specified requirements: fixing UI overlapping issues across all viewports, delivering a modern realistic EduAdmin theme, deploying a full-featured Application Tracking lifecycle timeline, and integrating an interactive simulated fee payment and electronic receipt generator. All code changes have been validated, tested in browser sessions, committed, and pushed to the remote GitHub repository.")

out_path = r"c:\WD project\docs\UniBridge_Project_Report.docx"
doc.save(out_path)
print(f"SUCCESS: Report successfully created at {out_path}")
